from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pandas as pd


MAX_FILE_SIZE = 20 * 1024 * 1024

SUPPORTED_TYPES: dict[str, set[str]] = {
    "df14a": {".xlsx", ".xls", ".csv"},
    "acta": {".pdf", ".docx"},
    "requisitos": {".xlsx", ".xls", ".csv"},
    "otro": {".xlsx", ".xls", ".csv", ".pdf", ".docx"},
}

TYPE_LABELS = {
    "df14a": "DF14A",
    "acta": "Acta de certificacion",
    "requisitos": "Requisitos / pendientes",
    "otro": "Otro archivo",
}


@dataclass
class ValidationResult:
    file_name: str
    information_type: str
    status: str
    checks: list[dict[str, str]]
    records_found: int | None = None
    valid_records: int | None = None
    inconsistencies: int = 0
    message: str | None = None

    def as_dict(self) -> dict:
        return {
            "file_name": self.file_name,
            "information_type": self.information_type,
            "status": self.status,
            "checks": self.checks,
            "records_found": self.records_found,
            "valid_records": self.valid_records,
            "inconsistencies": self.inconsistencies,
            "message": self.message,
        }


def validate_file(file_name: str, content: bytes, information_type: str) -> ValidationResult:
    checks: list[dict[str, str]] = []
    extension = Path(file_name).suffix.lower()
    label = TYPE_LABELS.get(information_type, information_type)

    if information_type not in SUPPORTED_TYPES:
        return _error(file_name, label, "Tipo de informacion no reconocido.")
    if not content:
        return _error(file_name, label, "El archivo esta vacio.")
    if len(content) > MAX_FILE_SIZE:
        return _error(file_name, label, "El archivo supera el limite recomendado de 20 MB.")

    checks.append({"label": "Archivo abierto correctamente", "status": "ok"})
    if extension not in SUPPORTED_TYPES[information_type]:
        return ValidationResult(
            file_name=file_name,
            information_type=label,
            status="error",
            checks=checks + [{"label": "Formato compatible", "status": "error"}],
            message=f"{extension or 'Sin extension'} no es compatible con {label}.",
        )

    checks.append({"label": "Formato compatible", "status": "ok"})
    records_found: int | None = None
    valid_records: int | None = None
    inconsistencies = 0
    try:
        if extension == ".csv":
            frame = pd.read_csv(BytesIO(content))
            checks.append({"label": "Columnas leidas correctamente", "status": "ok"})
            records_found, valid_records, inconsistencies, duplicate_rows = _tabular_summary(frame)
            if duplicate_rows:
                checks.append({"label": f"{duplicate_rows} registros duplicados detectados", "status": "warning"})
        elif extension in {".xlsx", ".xls"}:
            frame = pd.read_excel(BytesIO(content))
            checks.append({"label": "Columnas leidas correctamente", "status": "ok"})
            records_found, valid_records, inconsistencies, duplicate_rows = _tabular_summary(frame)
            if duplicate_rows:
                checks.append({"label": f"{duplicate_rows} registros duplicados detectados", "status": "warning"})
        elif extension == ".pdf":
            import fitz

            document = fitz.open(stream=content, filetype="pdf")
            if document.page_count == 0:
                raise ValueError("El PDF no contiene paginas.")
            checks.append({"label": "Documento PDF abierto correctamente", "status": "ok"})
        elif extension == ".docx":
            from docx import Document

            document = Document(BytesIO(content))
            if not document.paragraphs and not document.tables:
                raise ValueError("El documento no contiene contenido.")
            checks.append({"label": "Documento Word abierto correctamente", "status": "ok"})
    except Exception as error:
        return ValidationResult(
            file_name=file_name,
            information_type=label,
            status="warning",
            checks=checks + [{"label": "Contenido legible", "status": "warning"}],
            message=f"El formato parece correcto, pero no se pudo leer su contenido: {error}",
        )

    if records_found is not None:
        checks.append({"label": f"{records_found:,} registros encontrados".replace(",", "."), "status": "ok"})
        checks.append({"label": f"{valid_records:,} registros validos".replace(",", "."), "status": "ok"})
        if inconsistencies:
            checks.append({"label": f"{inconsistencies} registros presentan inconsistencias", "status": "warning"})

    return ValidationResult(
        file_name=file_name,
        information_type=label,
        status="warning" if inconsistencies else "ready",
        checks=checks,
        records_found=records_found,
        valid_records=valid_records,
        inconsistencies=inconsistencies,
        message="La validacion es preliminar y no se han guardado datos.",
    )


def _tabular_summary(frame: pd.DataFrame) -> tuple[int, int, int, int]:
    records_found = len(frame.index)
    invalid_rows = int(frame.isna().all(axis=1).sum())
    duplicate_rows = int(frame.duplicated(keep="first").sum())
    return records_found, records_found - invalid_rows, invalid_rows + duplicate_rows, duplicate_rows


def _error(file_name: str, label: str, message: str) -> ValidationResult:
    return ValidationResult(
        file_name=file_name,
        information_type=label,
        status="error",
        checks=[{"label": message, "status": "error"}],
        message=message,
    )