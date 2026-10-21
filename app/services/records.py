"""Interpretacion y preparacion de registros importables.

Este modulo concentra la decision de *que registros se pueden importar* a partir
de lo que devuelve cada lector de :mod:`app.services.parsers`.

Contadores usados en toda la aplicacion
---------------------------------------
``records_found``
    Registros que devolvio el lector del archivo. En ``actas_completo`` incluye
    todas las hojas procesadas (personas, complementaria y meta). Es el total
    fisico/logico detectado, no el numero de filas de una sola hoja.

``valid_records`` / ``importable_records``
    Registros que cumplen los campos obligatorios del tipo y que, despues de
    quitar duplicados dentro del mismo archivo, se guardarian en la base de
    datos. No incluye filas rechazadas ni duplicados internos.

``rejected_records``
    Registros detectados que no se pueden importar porque les falta un campo
    obligatorio o porque el contenido no es interpretable.

``duplicate_records``
    Registros validos repetidos dentro del mismo archivo. Se conserva el ultimo
    y se informa el resto; no se insertan filas duplicadas.

``new_records`` / ``updated_records`` (vista previa y confirmacion)
    Se calculan sobre los ``importable_records`` unicos, comparando su llave
    natural con la base de datos. Son exactamente las filas que se insertaran o
    actualizaran al confirmar.

Las filas con advertencias no bloqueantes (nivel sin clasificar, conflicto
CERTIFICADOS/PENDIENTES, fecha fuera de periodo) siguen siendo importables: la
advertencia queda visible para revision, pero no se descarta ni se modifica el
dato automaticamente.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd

from app.services import parsers


# Claves de la hoja META que representan metas reconocidas. ``fila_rotulos`` y
# ``*_celda`` son metadatos de trazabilidad y no cuentan como una meta.
META_VALUE_KEYS = {
    "superior",
    "tecnica",
    "articulacion",
    "complementaria",
    "total_titulada",
    "total_formacion_profesional",
    "campesena",
    "fullpopular",
}

_BLANK_VALUES = {"", "nan", "nat", "none", "null"}


@dataclass
class RecordEntry:
    information_type: str
    record: dict[str, Any]
    key: str
    index: int
    reason: str | None = None


@dataclass
class PreparedRecords:
    entries: list[RecordEntry] = field(default_factory=list)
    importable: list[RecordEntry] = field(default_factory=list)
    rejected: list[RecordEntry] = field(default_factory=list)
    duplicates: list[RecordEntry] = field(default_factory=list)

    @property
    def records_found(self) -> int:
        return len(self.entries)

    @property
    def valid_records(self) -> int:
        return len(self.importable)

    @property
    def rejected_records(self) -> int:
        return len(self.rejected)

    @property
    def duplicate_records(self) -> int:
        return len(self.duplicates)

    def breakdown(self) -> dict[str, dict[str, int]]:
        """Resumen por tipo interno (certificados, pendientes, meta, ...).

        ``valid_records`` cuenta los registros unicos que realmente se
        importarian. Asi la suma del desglose coincide con ``valid_records``
        aunque el archivo tenga filas duplicadas o rechazadas.
        """
        result: dict[str, dict[str, int]] = {}
        for entry in self.entries:
            bucket = result.setdefault(
                entry.information_type,
                {"records_found": 0, "valid_records": 0, "rejected_records": 0},
            )
            bucket["records_found"] += 1
            if entry.reason is not None:
                bucket["rejected_records"] += 1
        for entry in self.importable:
            bucket = result.setdefault(
                entry.information_type,
                {"records_found": 0, "valid_records": 0, "rejected_records": 0},
            )
            bucket["valid_records"] += 1
        return result


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return str(value).strip().lower() in _BLANK_VALUES


def _normalized(record: dict[str, Any]) -> dict[str, Any]:
    return {parsers.normalize_text(str(key)): value for key, value in record.items()}


def _first_present(record: dict[str, Any], *names: str) -> Any:
    normalized = _normalized(record)
    for name in names:
        value = normalized.get(parsers.normalize_text(name))
        if not _is_blank(value):
            return value
    return None


def _has_any_value(record: dict[str, Any]) -> bool:
    return any(not _is_blank(value) for value in record.values())


def importability_reason(information_type: str, record: dict[str, Any]) -> str | None:
    """Devuelve por que un registro no se puede importar, o ``None`` si es apto.

    Las reglas salen de ``REQUIRED_COLUMNS`` para los tipos tabulares y de la
    estructura propia de cada fuente para los demas. No se inventan requisitos
    administrativos: si una regla no esta documentada, el registro se conserva y
    se informa como advertencia en otro punto del flujo.
    """
    clean = {key: value for key, value in record.items() if key != "__type__"}

    if information_type in {"certificados", "pendientes"}:
        if _first_present(clean, "numero de documento", "documento") is None:
            return "Falta NUMERO DE DOCUMENTO"
        if _first_present(clean, "nombre", "nombres") is None:
            return "Falta NOMBRE"
        if _first_present(clean, "ficha", "codigo ficha") is None:
            return "Falta FICHA"
        return None

    if information_type == "df14a":
        if _first_present(clean, "ficha", "codigo ficha") is None:
            return "Falta FICHA"
        return None

    if information_type == "requisitos":
        if _first_present(clean, "documento", "numero de documento", "identificacion") is None:
            return "Falta DOCUMENTO"
        if _first_present(clean, "nombres", "nombre") is None:
            return "Falta NOMBRES"
        if _first_present(clean, "requisito", "tipo requisito") is None:
            return "Falta REQUISITO"
        if _first_present(clean, "estado") is None:
            return "Falta ESTADO"
        return None

    if information_type == "complementaria":
        cantidad = clean.get("cantidad")
        if not isinstance(cantidad, int) or isinstance(cantidad, bool):
            return "Cantidad invalida"
        return None

    if information_type == "meta":
        if not any(key in clean for key in META_VALUE_KEYS):
            return "No se reconocieron metas"
        return None

    if information_type == "informe":
        if _first_present(clean, "indicador") is None:
            return "Fila sin indicador"
        return None

    if information_type == "acta":
        # Un acta puede aportar filas de tabla o, si no tiene tabla, un registro
        # de metadatos del documento. En ambos casos debe tener contenido real.
        if not _has_any_value(clean):
            return "Registro de acta vacio"
        return None

    if not _has_any_value(clean):
        return "Registro vacio"
    return None


def prepare_records(
    file_name: str,
    content: bytes,
    information_type: str = "otro",
) -> PreparedRecords:
    """Lee el archivo y separa registros importables, rechazados y duplicados."""
    raw = parse_records(file_name, content, information_type)
    prepared = PreparedRecords()
    unique: dict[tuple[str, str], RecordEntry] = {}

    for index, raw_record in enumerate(raw):
        record = dict(raw_record)
        rtype = str(record.pop("__type__", information_type))
        key = record_key(rtype, record)
        reason = importability_reason(rtype, record)
        entry = RecordEntry(
            information_type=rtype,
            record=record,
            key=key,
            index=index,
            reason=reason,
        )
        prepared.entries.append(entry)
        if reason is not None:
            prepared.rejected.append(entry)
            continue
        unique_key = (rtype, key)
        if unique_key in unique:
            # Misma llave natural dentro del archivo: se conserva el ultimo
            # valor (igual que una actualizacion) y se informa el duplicado.
            prepared.duplicates.append(entry)
            unique[unique_key] = entry
        else:
            unique[unique_key] = entry

    prepared.importable = list(unique.values())
    return prepared


def parse_records(
    file_name: str,
    content: bytes,
    information_type: str = "otro",
) -> list[dict[str, Any]]:
    """Convierte un archivo en registros, usando el lector de su tipo.

    La decision depende del tipo de informacion y de la extension. Los lectores
    de :mod:`app.services.parsers` ya devuelven nombres de campo internos.
    """
    extension = Path(file_name).suffix.lower()

    if information_type == "df14a":
        if extension == ".xml":
            frame = parsers.parse_df14a_xml(content)
            return _clean_identifier_fields(parsers.frame_to_records(frame))
        if extension in {".csv", ".xlsx", ".xls"}:
            return _clean_identifier_fields(_parse_generic_tabular(extension, content))
        return []
    if information_type == "certificados" and extension in {".xlsx", ".xls"}:
        return parsers.read_certificados_records(content)
    if information_type == "pendientes" and extension in {".xlsx", ".xls"}:
        return parsers.read_pendientes_records(content)
    if information_type in {"certificados", "pendientes"} and extension == ".csv":
        return parsers.read_actas_csv_records(content)
    if information_type == "csv_certificacion" and extension == ".csv":
        frame = parsers.read_csv_frame(content)
        sub_type, _ = parsers.detect_csv_certificacion_type(frame)
        records = parsers.read_actas_csv_records(content)
        for record in records:
            record["__type__"] = sub_type
        return records
    if information_type == "complementaria" and extension in {".xlsx", ".xls"}:
        return parsers.read_complementaria_records(content)
    if information_type == "meta" and extension in {".xlsx", ".xls"}:
        metas = parsers.read_meta_sheet(content)
        return [metas] if metas else []
    if information_type == "actas_completo" and extension in {".xlsx", ".xls"}:
        return parsers.read_actas_full(content)
    if information_type == "informe" and extension == ".docx":
        return parsers.parse_informe_referencia(content)
    if information_type == "requisitos":
        if extension in {".xlsx", ".xls", ".csv"}:
            return _parse_generic_tabular(extension, content)
        return []
    if information_type == "acta":
        if extension == ".docx":
            result = parsers.parse_acta_docx(content)
            records = result.get("records", [])
            if records:
                return records
            if result.get("paragraphs"):
                return [_document_metadata(file_name, extension, result.get("paragraphs"))]
            return []
        if extension == ".pdf":
            return [_pdf_metadata(file_name, content)]
        return []
    if extension in {".pdf", ".docx"}:
        # Documentos de referencia ("otro"): se conserva un registro de
        # metadatos, pero solo si el archivo pudo abrirse y tiene contenido. Un
        # PDF/DOCX corrupto devuelve lista vacia y la validacion lo rechaza
        # antes de confirmar.
        try:
            if extension == ".pdf":
                _pdf_page_count(content)
                return [{"file_name": file_name, "format": "pdf"}]
            return [_docx_metadata(file_name, content)]
        except Exception:
            return []
    return _parse_generic_tabular(extension, content)


# Columnas cuyo valor es un identificador. Excel puede devolverlas como decimal
# (``1043966273.0``) o dejar la celda vacia, y el mismo aprendiz llegaria con dos
# valores distintos segun el formato de origen. Se normalizan con
# ``clean_identification`` para que la llave natural y el dato guardado coincidan.
_IDENTIFIER_FIELDS = {
    "ficha",
    "codigo ficha",
    "documento",
    "numero de documento",
    "numero documento",
    "identificacion",
    "nis",
    "acta",
    "numero de acta",
}


def _clean_identifier_fields(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Normaliza documento, ficha y acta de cada registro, sin renombrar columnas."""
    cleaned: list[dict[str, Any]] = []
    for record in records:
        nuevo: dict[str, Any] = {}
        for key, value in record.items():
            if parsers.normalize_text(str(key)) in _IDENTIFIER_FIELDS:
                value = parsers.clean_identification(value)
                if value is None:
                    continue
            nuevo[key] = value
        cleaned.append(nuevo)
    return cleaned


def _parse_generic_tabular(extension: str, content: bytes) -> list[dict[str, Any]]:
    if extension == ".csv":
        frame = parsers.read_csv_frame(content)
        return _frame_records(frame)
    if extension in {".xlsx", ".xls"}:
        frame = pd.read_excel(BytesIO(content))
        return _frame_records(frame)
    return []


def _frame_records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    clean_frame = frame.where(pd.notna(frame), None)
    return [
        {str(key): value for key, value in row.items() if value is not None}
        for row in clean_frame.to_dict(orient="records")
        if any(value is not None for value in row.values())
    ]


def _document_metadata(
    file_name: str,
    extension: str,
    paragraphs: list[str] | None = None,
) -> dict[str, Any]:
    """Registro de metadatos para un DOCX valido que no aporta filas de tabla."""
    return {
        "file_name": file_name,
        "format": extension.lstrip("."),
        "paragraphs": len(paragraphs or []),
    }


def _docx_metadata(file_name: str, content: bytes) -> dict[str, Any]:
    """Metadatos de un DOCX legible que no se interpreta como tabla."""
    from docx import Document

    document = Document(BytesIO(content))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if not paragraphs and not document.tables:
        raise ValueError("El documento no contiene texto ni tablas interpretables.")
    return {
        "file_name": file_name,
        "format": "docx",
        "paragraphs": len(paragraphs),
        "tables": len(document.tables),
    }


def _pdf_page_count(content: bytes) -> int:
    """Comprueba que el PDF se abra y tenga al menos una pagina."""
    import fitz

    document = fitz.open(stream=content, filetype="pdf")
    try:
        if document.page_count == 0:
            raise ValueError("El PDF no contiene paginas.")
        document.load_page(0)
        return document.page_count
    finally:
        document.close()


def _pdf_metadata(file_name: str, content: bytes) -> dict[str, Any]:
    """Registro de metadatos de un PDF valido.

    No se inventan aprendices ni filas: solo se deja constancia del documento.
    Si el PDF no se puede abrir, la excepcion se propaga para que la validacion
    lo rechace y no se genere ningun registro ficticio.
    """
    import fitz

    document = fitz.open(stream=content, filetype="pdf")
    try:
        if document.page_count == 0:
            raise ValueError("El PDF no contiene paginas.")
        pages = document.page_count
        first_page = document.load_page(0)
        text_preview = " ".join(first_page.get_text("text").split())[:200]
    finally:
        document.close()
    return {
        "file_name": file_name,
        "format": "pdf",
        "pages": pages,
        "text_preview": text_preview,
    }


def record_key(information_type: str, record: dict[str, Any]) -> str:
    """Llave natural usada para no duplicar registros entre importaciones."""
    clean = {k: v for k, v in record.items() if k != "__type__"}

    if information_type in {"certificados", "pendientes"}:
        doc = parsers.clean_identification(
            _first_present(clean, "numero de documento", "documento")
        )
        if doc:
            ficha = (
                parsers.clean_identification(_first_present(clean, "ficha", "codigo ficha"))
                or ""
            )
            return hashlib.sha256(
                f"{information_type}:{doc}:{ficha}".encode("utf-8")
            ).hexdigest()

    if information_type == "df14a":
        doc = parsers.clean_identification(
            _first_present(
                clean,
                "numero de documento",
                "documento",
                "identificacion",
                "nis",
            )
        )
        if doc:
            ficha = (
                parsers.clean_identification(
                    _first_present(clean, "ficha", "codigo ficha")
                )
                or ""
            )
            return hashlib.sha256(
                f"{information_type}:{doc}:{ficha}".encode("utf-8")
            ).hexdigest()

    if information_type == "requisitos":
        doc = parsers.clean_identification(
            _first_present(clean, "documento", "numero de documento", "identificacion")
        )
        req = _first_present(clean, "requisito", "tipo requisito")
        if doc and req:
            return hashlib.sha256(
                f"{information_type}:{doc}:{parsers.normalize_text(req)}".encode("utf-8")
            ).hexdigest()

    normalized = json.dumps(
        {"information_type": information_type, "record": clean},
        ensure_ascii=True,
        sort_keys=True,
        default=str,
    )
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()
